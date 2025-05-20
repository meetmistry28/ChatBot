import csv
import re
from docx import Document  


def clean_text(text):
    text = text.lower()
    text = re.sub(r"[^a-zA-Z0-9\s]", "", text)
    words = text.split()
    return words


def get_common_words(words1, words2):
    words1_set = set(words1)
    words2_set = set(words2)
    common_words = words1_set.intersection(words2_set)
    return common_words


def merge_data(file1, file2, output_csv, output_docx="merged_output.docx"):
    with open(file1, mode="r", encoding="utf-8") as f1:
        reader1 = csv.DictReader(f1)
        data1 = [row for row in reader1]
    with open(file2, mode="r", encoding="utf-8") as f2:
        reader2 = csv.DictReader(f2)
        data2 = [row for row in reader2]

    merged_data = []

    # Initialize Word document
    doc = Document()
    doc.add_heading("Merged Output", 0)

    for row1 in data1:
        title1 = row1.get("Key", "").strip()
        paragraph1 = row1.get("Value", "").strip()
        words1 = clean_text(title1) + clean_text(paragraph1)

        for row2 in data2:
            title2 = row2.get("Title", "").strip()
            bio2 = row2.get("Bio", "").strip()
            words2 = clean_text(title2) + clean_text(bio2)

            common_words = get_common_words(words1, words2)

            if common_words:
                merged_row = {
                    "Title Wiki": title1,
                    "Paragraph Wiki": paragraph1,
                    "Title BMS": title2,
                    "Bio BMS": bio2,
                    # "Common Words": ", ".join(common_words),
                }
                merged_data.append(merged_row)

                # ✅ Write to Word
                doc.add_heading("Match", level=1)
                doc.add_paragraph(f"Title Wiki: {title1}")
                doc.add_paragraph(f"Paragraph Wiki: {paragraph1}")
                doc.add_paragraph(f"Title BMS: {title2}")
                doc.add_paragraph(f"Bio BMS: {bio2}")
                # doc.add_paragraph(f"Common Words: {', '.join(common_words)}")
                doc.add_paragraph("\n")

    # Save to CSV
    with open(output_csv, mode="w", newline="", encoding="utf-8") as out_file:
        fieldnames = [
            "Title Wiki",
            "Paragraph Wiki",
            "Title BMS",
            "Bio BMS",
            # "Common Words",
        ]
        writer = csv.DictWriter(out_file, fieldnames=fieldnames)
        writer.writeheader()
        for row in merged_data:
            writer.writerow(row)

    # Save to Word
    doc.save(output_docx)
    print(f"✅ Merged data saved in {output_csv} and {output_docx}")


# Call the function
merge_data("wiki_output_clean1.csv", "bms_output.csv", "merged_output.csv")
