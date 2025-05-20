from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from bs4 import BeautifulSoup
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import csv
from docx import Document

class WikiScraper:
    def __init__(self, driver_path):
        service = Service(driver_path)
        options = webdriver.ChromeOptions()
    
        options.add_argument("--disable-blink-features=AutomationControlled")
        options.add_argument("user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                             "AppleWebKit/537.36 (KHTML, like Gecko) "
                             "Chrome/122.0.0.0 Safari/537.36")
        options.add_experimental_option("excludeSwitches", ["enable-automation"])
        options.add_experimental_option('useAutomationExtension', False)

        self.driver = webdriver.Chrome(service=service, options=options)

        self.driver.execute_cdp_cmd("Page.addScriptToEvaluateOnNewDocument", {
            "source": """
            Object.defineProperty(navigator, 'webdriver', {
                get: () => undefined
            })
            """
        })


    def scrape_page(self, url):
        self.driver.get(url)

        try:
            WebDriverWait(self.driver, 10).until(
                EC.presence_of_element_located((By.TAG_NAME, "h1"))
            )
        except:
            print(f"⚠️ Timeout waiting for page to load: {url}")

        soup = BeautifulSoup(self.driver.page_source, "html.parser")
        data = {}

        title = soup.find("div", {"class": "sc-122x28k-4 gNXQur"})
        data["Title"] = title.text.strip() if title else "N/A"

        paragraph = "N/A"
        try:
            about_sections = soup.find_all("div", string="About")
            for about in about_sections:
                sibling = about.find_next_sibling("div")
                if sibling:
                    paragraph = sibling.get_text(strip=True)
                    break
        except Exception as e:
            paragraph = f"Error: {e}"

        data["Bio"] = paragraph
        data["URL"] = url
        return data

    def scrape_multiple(self, urls):
        all_data = []
        for url in urls:
            page_data = self.scrape_page(url)
            all_data.append(page_data)
        return all_data

    def save_to_word(self, all_data, filename="bms_output.docx"):
        doc = Document()
        doc.add_heading("BookMyShow Scraped Data", level=1)

        for entry in all_data:
            doc.add_heading(entry.get("Title", "N/A"), level=2)
            doc.add_paragraph(f"Bio: {entry.get('Bio', 'N/A')}")
            doc.add_paragraph(f"URL: {entry.get('URL', 'N/A')}")
            doc.add_paragraph("\n")

        doc.save(filename)

    def close(self):
        self.driver.quit()


if __name__ == "__main__":
    urls = [
        "https://in.bookmyshow.com/person/rajinikanth/1795",
    ]

    scraper = WikiScraper(
        driver_path=r"D:\helloworld\chromedriver-win64\chromedriver.exe"
    )
    all_data = scraper.scrape_multiple(urls)
    scraper.save_to_word(all_data)
    scraper.close()
    print("✅ Data saved in Word format.")
