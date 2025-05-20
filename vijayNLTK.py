import os, re, nltk
from collections import defaultdict
from nltk.tokenize import word_tokenize, sent_tokenize
from nltk.corpus import stopwords
from nltk.stem import WordNetLemmatizer
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

nltk.download(['punkt', 'stopwords', 'wordnet'])

class DOCXChatbot:
    def __init__(self, docx_path):
        self.docx_path = docx_path
        self.lemmatizer = WordNetLemmatizer()
        self.stop_words = set(stopwords.words('english'))
        self.document_text = self._load_docx()
        self.sentences = sent_tokenize(self.document_text)
        self.qa_pairs = self._extract_qa_pairs()
        self.answer_bank, self.answer_pos = defaultdict(list), defaultdict(int)

    def _is_bold(self, para):
        return any(run.bold for run in para.runs)

    def _load_docx(self):
        try:
            return '\n'.join(
                p.text.strip() for p in Document(self.docx_path).paragraphs
                if p.text.strip() and not self._is_bold(p)
            )
        except: return ""

    def _extract_qa_pairs(self):
        pairs, q = defaultdict(list), None
        for p in Document(self.docx_path).paragraphs:
            t = p.text.strip()
            if not t or self._is_bold(p): continue
            if t.endswith('?'): q = t
            elif q: pairs[q].append(t)
        return pairs

    def _preprocess(self, text):
        text = re.sub(r'[^\w\s]|\d+', '', text.lower())
        return ' '.join(self.lemmatizer.lemmatize(w) for w in word_tokenize(text) if w not in self.stop_words and len(w) > 2)

    def _similarity(self, t1, t2):
        texts = [self._preprocess(t1), self._preprocess(t2)]
        try:
            tfidf = TfidfVectorizer().fit_transform(texts)
            return cosine_similarity(tfidf[0:1], tfidf[1:2])[0][0]
        except: return 0.0

    def get_answer(self, question, threshold=0.2):
        if not question.strip(): return "Please ask a question about Vijay Deverakonda."

        if question not in self.answer_bank:
            qa = [a for q, a_list in self.qa_pairs.items() if self._similarity(question, q) >= threshold for a in a_list]
            sents = [s for s in self.sentences if self._similarity(question, s) >= threshold]
            self.answer_bank[question] = qa + sents or ["I don't have an answer for that.", "That information isn't available."]

        answers = self.answer_bank[question]
        pos = self.answer_pos[question]
        answer = answers[pos % len(answers)].strip()
        self.answer_pos[question] += 1
        return answer

if __name__ == "__main__":
    docx_file = "Vijay.docx"
    if not os.path.exists(docx_file):
        print(f"Error: File '{docx_file}' not found."); exit()

    bot = DOCXChatbot(docx_file)
    print("Chatbot ready! (type 'exit', 'quit', or 'bye' to stop)\n")
    while True:
        user_input = input("You: ").strip()
        if user_input.lower() in ['quit', 'exit', 'bye']:
            print("Bot: Goodbye!"); break
        print("Bot:", bot.get_answer(user_input))