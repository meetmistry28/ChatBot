from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from bs4 import BeautifulSoup
from docx import Document
import time


class WikiScraper:
    def __init__(self, driver_path):
        service = Service(driver_path)
        options = webdriver.ChromeOptions()
        options.add_argument("--headless")
        self.driver = webdriver.Chrome(service=service, options=options)

    def scrape_page(self, url):
        self.driver.get(url)
        time.sleep(3)
        soup = BeautifulSoup(self.driver.page_source, "html.parser")
        data = {}
        try:
            title = soup.find("h1", {"id": "firstHeading"}).text.strip()
        except:
            title = "N/A"
        data["title"] = title  
        try:
            content_div = soup.find("div", {"class": "vector-body ve-init-mw-desktopArticleTarget-targetContainer"})
            paragraphs = content_div.find_all("p")
            paragraphs_text = []
            for p in paragraphs:
                text = p.get_text(strip=True)
                if text:
                    paragraphs_text.append(text)
            paragraph = "\n\n".join(paragraphs_text)
        except Exception as e:
            paragraph = f"Error: {e}"
        data["paragraph"] = paragraph
        infobox = soup.find("table", {"class": "infobox biography vcard"})
        if infobox:
            rows = infobox.find_all("tr")
            for row in rows:
                header = row.find("th")
                value = row.find("td")
                if header and value:
                    label = header.text.strip().lower()
                    content = value.text.strip().replace("\n", " ")
                    data[label] = content
                    if label == "born":
                        break  
        data["url"] = url
        return data

    def scrape_multiple(self, urls):
        all_data = []
        for url in urls:
            page_data = self.scrape_page(url)
            for key, value in page_data.items():
                all_data.append({"Key": key, "Value": value})
            all_data.append({"Key": "", "Value": ""})
        return all_data

    def save_to_word(self, all_data, filename="LanaDelRey.docx"):
        doc = Document()
        doc.add_heading("Wikipedia Scraped Data", level=1)

        for row in all_data:
            key = row.get("Key", "")
            value = row.get("Value", "")
            if key == "" and value == "":
                doc.add_paragraph("\n")  # Add space between entries
            else:
                doc.add_paragraph(f"{key.title()}: {value}")

        doc.save(filename)

    def close(self):
        self.driver.quit()


if __name__ == "__main__":
    urls = [
        "https://en.wikipedia.org/wiki/Lana_Del_Rey",
    ]

    scraper = WikiScraper(
        driver_path=r"D:\helloworld\chromedriver-win64\chromedriver.exe"
    )
    all_data = scraper.scrape_multiple(urls)
    scraper.save_to_word(all_data)
    scraper.close()
    print("✅ Data saved in Word format.")

# from selenium import webdriver
# from selenium.webdriver.chrome.service import Service
# from bs4 import BeautifulSoup
# from docx import Document
# import time

# class WikiScraper:
#     def __init__(self, driver_path):
#         service = Service(driver_path)
#         options = webdriver.ChromeOptions()
#         options.add_argument("--headless")
#         self.driver = webdriver.Chrome(service=service, options=options)

#     def scrape_page(self, url):
#         self.driver.get(url)
#         time.sleep(3)
#         soup = BeautifulSoup(self.driver.page_source, "html.parser")
#         data = {}
        
#         # Extract title
#         try:
#             title = soup.find("h1", {"id": "firstHeading"}).text.strip()
#         except:
#             title = "N/A"
#         data["title"] = title  

#         # Extract paragraphs
#         try:
#             content_div = soup.find("div", {"class": "vector-body ve-init-mw-desktopArticleTarget-targetContainer"})
#             paragraphs = content_div.find_all("p")
#             paragraphs_text = []
#             for p in paragraphs:
#                 text = p.get_text(strip=True)
#                 if text:
#                     paragraphs_text.append(text)
#             paragraph = "\n\n".join(paragraphs_text)
#         except Exception as e:
#             paragraph = f"Error: {e}"
#         data["paragraph"] = paragraph

#         # Extract all <h2> elements
#         try:
#             h2_elements = soup.find_all("h2")
#             h2_texts = [h2.get_text(strip=True) for h2 in h2_elements if h2.get_text(strip=True)]
#             data["headings"] = "\n".join(h2_texts) if h2_texts else "No <h2> headings found"
#         except Exception as e:
#             data["headings"] = f"Error extracting <h2>: {e}"

#         # Extract infobox
#         infobox = soup.find("table", {"class": "infobox biography vcard"})
#         if infobox:
#             rows = infobox.find_all("tr")
#             for row in rows:
#                 header = row.find("th")
#                 value = row.find("td")
#                 if header and value:
#                     label = header.text.strip().lower()
#                     content = value.text.strip().replace("\n", " ")
#                     data[label] = content
#                     if label == "born":
#                         break  
#         data["url"] = url
#         return data

#     def scrape_multiple(self, urls):
#         all_data = []
#         for url in urls:
#             page_data = self.scrape_page(url)
#             for key, value in page_data.items():
#                 all_data.append({"Key": key, "Value": value})
#             all_data.append({"Key": "", "Value": ""})
#         return all_data

#     def save_to_word(self, all_data, filename="wiki_output_clean1.docx"):
#         doc = Document()
#         doc.add_heading("Wikipedia Scraped Data", level=1)

#         for row in all_data:
#             key = row.get("Key", "")
#             value = row.get("Value", "")
#             if key == "" and value == "":
#                 doc.add_paragraph("\n")  # Add space between entries
#             else:
#                 doc.add_paragraph(f"{key.title()}: {value}")

#         doc.save(filename)

#     def close(self):
#         self.driver.quit()

# if __name__ == "__main__":
#     urls = [
#         "https://en.wikipedia.org/wiki/Rajinikanth",
#     ]

#     scraper = WikiScraper(
#         driver_path=r"D:\helloworld\chromedriver-win64\chromedriver.exe"
#     )
#     all_data = scraper.scrape_multiple(urls)
#     scraper.save_to_word(all_data)
#     scraper.close()
#     print("✅ Data saved in Word format.")